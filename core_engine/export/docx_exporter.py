from pathlib import Path
from typing import Any, Dict, List, Mapping, Union
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches


Paragraph = Dict[str, Any]
ParagraphList = List[Paragraph]
PathLike = Union[str, Path]


# ================================
#       СТИЛИ ТЕКСТА
# ================================

HEADING_STYLES: Mapping[str, str] = {
    "heading1": "Heading 1",
    "heading2": "Heading 2",
    "heading": "Heading 2",
    "h1": "Heading 1",
    "h2": "Heading 2",
}

LIST_TYPES = {"list_item"}
CAPTION_TYPES = {"caption"}
FIGURE_TYPES = {"figure"}
TABLE_TYPES = {"table"}


# ================================
#           УТИЛИТЫ
# ================================

def _ensure_path(path: PathLike) -> Path:
    p = Path(path)
    if not p.parent.exists():
        p.parent.mkdir(parents=True, exist_ok=True)
    return p


def _coerce_paragraph_list(paragraphs: Any) -> ParagraphList:
    if not isinstance(paragraphs, list):
        raise TypeError(f"EXPORT ERROR: paragraphs must be list, got {type(paragraphs)}")

    cleaned = [p for p in paragraphs if isinstance(p, dict)]
    if not cleaned:
        raise ValueError("EXPORT ERROR: no valid paragraph dicts in input")

    return cleaned


def _validate_paragraphs(paragraphs: ParagraphList) -> None:
    if not paragraphs:
        raise ValueError("EXPORT ERROR: paragraph_stream is empty")

    if all(not (p.get("text") or "").strip() for p in paragraphs):
        raise ValueError("EXPORT ERROR: all paragraphs empty")


def _get_para_type(p: Paragraph) -> str:
    t = p.get("type") or ""
    return t.lower().strip() if isinstance(t, str) else ""


def _get_para_text(p: Paragraph) -> str:
    txt = p.get("text")
    if txt is None:
        return ""
    if not isinstance(txt, str):
        txt = str(txt)
    return txt.replace("\r\n", "\n").replace("\r", "\n").strip()


# ================================
#         РЕНДЕР БЛОКОВ
# ================================

def _add_heading(doc: Document, text: str, level: int = 2) -> None:
    para = doc.add_heading(text, level=level)
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    # Улучшенные стили для заголовков
    para_format = para.paragraph_format
    if level == 1:
        para_format.space_before = Pt(12)
        para_format.space_after = Pt(6)
    elif level == 2:
        para_format.space_before = Pt(10)
        para_format.space_after = Pt(4)
    else:
        para_format.space_before = Pt(8)
        para_format.space_after = Pt(4)


def _add_list_item(doc: Document, text: str) -> None:
    """
    Улучшенная обработка списков:
    - Нумерованные списки (1. 2. 3. или 1) 2) 3))
    - Маркированные списки (• - *)
    """
    text_stripped = text.lstrip()
    
    # Нумерованный список
    numbered_match = re.match(r"^\(?\d+[\.\)]\s+", text_stripped)
    if numbered_match:
        para = doc.add_paragraph(style="List Number")
        # Убираем номер из текста (Word сам пронумерует)
        text_clean = re.sub(r"^\(?\d+[\.\)]\s+", "", text_stripped)
        run = para.add_run(text_clean)
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    # Римские цифры
    elif re.match(r"^[ivxlcdm]+\.\s+", text_stripped.lower()):
        para = doc.add_paragraph(style="List Number")
        text_clean = re.sub(r"^[ivxlcdm]+\.\s+", "", text_stripped, flags=re.IGNORECASE)
        run = para.add_run(text_clean)
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    # Маркированный список
    else:
        para = doc.add_paragraph(style="List Bullet")
        # Убираем маркер если есть
        text_clean = re.sub(r"^[•\-\–\—\*]\s+", "", text_stripped)
        run = para.add_run(text_clean)
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    
    run.bold = False


def _add_caption(doc: Document, text: str) -> None:
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.italic = True
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _add_figure(doc: Document, text: str) -> None:
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold = True
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER


# ================================
#         ТАБЛИЦЫ DOCX v7.2
# ================================

def _add_table(doc: Document, rows: List[List[str]]) -> None:
    """
    Полноценная native-таблица DOCX v7.2:
    - создаёт таблицу
    - первая строка → header bold
    - fallback при ошибке
    """
    if not rows or not any(rows):
        para = doc.add_paragraph("[TABLE EMPTY]")
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        return

    # Определяем число колонок по максимальной длине строки
    max_cols = max(len(r) for r in rows)

    table = doc.add_table(rows=len(rows), cols=max_cols)
    table.style = "Table Grid"

    for r_idx, row in enumerate(rows):
        for c_idx in range(max_cols):
            cell = table.rows[r_idx].cells[c_idx]
            if c_idx < len(row):
                text = row[c_idx]
            else:
                text = ""

            run = cell.paragraphs[0].add_run(text)

            if r_idx == 0:
                run.bold = True

    # Добавляем отступ после таблицы
    doc.add_paragraph("")


def _add_table_placeholder(doc: Document, text: str) -> None:
    para = doc.add_paragraph()
    run = para.add_run(text)
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT


def _add_normal_paragraph(doc: Document, text: str) -> None:
    para = doc.add_paragraph(text)
    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    # Улучшенные стили: отступы и интервалы
    para_format = para.paragraph_format
    para_format.space_after = Pt(6)  # Небольшой отступ после параграфа
    para_format.first_line_indent = Inches(0)  # Без красной строки (как в научных статьях)


# ================================
#             EXPORT
# ================================

def export_docx(paragraphs: Any, output_path: PathLike) -> str:
    """
    paragraphs: список dict:
        {
            "type": "...",
            "text": "...",
            "rows": [...],  # если таблица
            ...
        }
    """

    output = _ensure_path(output_path)

    para_list = _coerce_paragraph_list(paragraphs)
    _validate_paragraphs(para_list)

    # Жёсткий контроль: если все тексты пустые — останавливаем экспорт
    non_empty = [p for p in para_list if (p.get("text") or "").strip()]
    if not non_empty:
        raise ValueError("EXPORT ERROR: all paragraphs are empty")

    doc = Document()
    prev_page = 0

    for p in para_list:
        p_type = _get_para_type(p)
        text = _get_para_text(p)
        page = p.get("page", 0)

        # ---------------------
        #      Page Break
        # ---------------------
        if p_type == "page_break":
            doc.add_page_break()
            prev_page = page
            continue

        if not text and p_type not in TABLE_TYPES:
            continue

        # ---------------------
        #      Заголовки
        # ---------------------
        if p_type in {"heading1", "h1"}:
            _add_heading(doc, text, level=1)
            prev_page = page
            continue

        if p_type in {"heading2", "heading", "h2"}:
            _add_heading(doc, text, level=2)
            prev_page = page
            continue

        if p_type in {"heading3", "h3"}:
            _add_heading(doc, text, level=3)
            prev_page = page
            continue

        # ---------------------
        #         Списки
        # ---------------------
        if p_type in LIST_TYPES:
            _add_list_item(doc, text)
            continue

        # ---------------------
        #     Подписи / Фигуры
        # ---------------------
        if p_type in CAPTION_TYPES:
            _add_caption(doc, text)
            continue

        if p_type in FIGURE_TYPES:
            _add_figure(doc, text)
            continue

        # ---------------------
        #         Таблицы
        # ---------------------
        if p_type in TABLE_TYPES:
            rows = p.get("rows") or []
            try:
                _add_table(doc, rows)
            except Exception:
                _add_table_placeholder(doc, text or "[TABLE]")
            continue

        # ---------------------
        #   Обычный параграф
        # ---------------------
        _add_normal_paragraph(doc, text)
        prev_page = page

    if len(doc.paragraphs) == 0:
        raise ValueError("EXPORT ERROR: empty DOCX")

    doc.save(str(output))
    return str(output)
